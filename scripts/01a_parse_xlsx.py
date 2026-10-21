#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
01a_parse_xlsx.py
=================
Parse the FinancialStatement .xlsx workbooks downloaded by 00_download.py into
one firm-year table (../data/panel_xlsx.csv).

IDX statements follow three balance-sheet taxonomies (general = 12xxxxx,
property = 22xxxxx, infrastructure = 32xxxxx), each with a classified and an
unclassified variant; income-statement sheets use the matching 13/23/33
prefixes and, in some filings, changed sheet codes across reporting years. The
parser scans all candidate sheets and extracts the items needed to build the
Altman Z"-Score components and the additional ratios, plus a prior-year net
income used by the persistent-distress sensitivity label.

Inputs:   downloads/<year>/<CODE>.xlsx
Output:   ../data/panel_xlsx.csv
Usage:    pip install openpyxl ; python 01a_parse_xlsx.py
"""

import csv
from pathlib import Path
from openpyxl import load_workbook

YEARS = [2021, 2022, 2023, 2024, 2025]
HERE = Path(__file__).resolve().parent
IN_DIR = HERE / "downloads"
OUT = HERE / ".." / "data" / "panel_xlsx.csv"

SHEET_INFO = "1000000"
SHEETS_NERACA = ["1210000", "2210000", "3210000"]          # classified
SHEETS_NERACA_UNCLASS = ["1220000", "2220000", "3220000"]  # unclassified
PREFIX_LABARUGI = ("13", "23", "33")

NERACA_LABELS = {
    "aset_lancar": "Jumlah aset lancar",
    "total_aset": "Jumlah aset",
    "liab_lancar": "Jumlah liabilitas jangka pendek",
    "total_liab": "Jumlah liabilitas",
    "saldo_laba_ditentukan": "Saldo laba yang telah ditentukan penggunaannya",
    "saldo_laba_belum": "Saldo laba yang belum ditentukan penggunaannya",
    "total_ekuitas": "Jumlah ekuitas",
}
LABARUGI_LABELS = {
    "pendapatan": "Penjualan dan pendapatan usaha",
    "laba_sebelum_pajak": "Jumlah laba (rugi) sebelum pajak",
    "beban_bunga": "Beban bunga dan keuangan",
    "beban_keuangan_alt": "Beban keuangan",
    "laba_bersih": "Jumlah laba (rugi)",
}
INFO_LABELS = {
    "nama": "Nama entitas", "industri_utama": "Industri utama entitas",
    "sektor": "Sektor", "subsektor": "Subsektor",
    "mata_uang": "Mata uang pelaporan",
    "pembulatan": "Pembulatan yang digunakan", "opini_audit": "Jenis opini auditor",
}


def first_number(cells):
    for v in cells:
        if isinstance(v, (int, float)):
            return float(v)
    return None


def first_text(cells):
    for v in cells:
        if v is not None and str(v).strip() != "":
            return str(v).strip()
    return None


def scan_sheet(ws, wanted, numeric=True):
    out = {k: None for k in wanted}
    order = sorted(wanted.items(), key=lambda kv: -len(kv[1]))
    for row in ws.iter_rows(values_only=True):
        vals = [v for v in row if v is not None]
        if not vals:
            continue
        label = str(vals[0]).strip()
        for key, prefix in order:
            if out[key] is None and label.startswith(prefix):
                if prefix == "Jumlah aset" and ("lancar" in label or "tidak lancar" in label):
                    continue
                if (prefix == "Jumlah liabilitas"
                        and ("jangka" in label or "dan ekuitas" in label) and key == "total_liab"):
                    continue
                if prefix == "Jumlah ekuitas" and "diatribusikan" in label:
                    continue
                if (prefix == "Jumlah laba (rugi)" and key == "laba_bersih"
                        and ("sebelum" in label or "operasi" in label or "per saham" in label)):
                    continue
                rest = vals[1:]
                out[key] = first_number(rest) if numeric else first_text(rest)
                break
    return out


def scan_two_numbers(ws, prefix, guards=()):
    for row in ws.iter_rows(values_only=True):
        vals = [v for v in row if v is not None]
        if not vals:
            continue
        label = str(vals[0]).strip()
        if label.startswith(prefix) and not any(g in label for g in guards):
            nums = [v for v in vals[1:] if isinstance(v, (int, float))]
            return (float(nums[0]) if nums else None,
                    float(nums[1]) if len(nums) > 1 else None)
    return (None, None)


def parse_file(path, year):
    row = {"kode": path.stem, "tahun": year}
    flags = []
    try:
        wb = load_workbook(path, read_only=True)
    except Exception as e:
        row["flag_masalah"] = f"GAGAL_BUKA: {e}"
        return row
    sheets = wb.sheetnames

    if SHEET_INFO in sheets:
        row.update(scan_sheet(wb[SHEET_INFO], INFO_LABELS, numeric=False))
    else:
        flags.append("TANPA_SHEET_INFO")

    bs = next((s for s in SHEETS_NERACA if s in sheets), None)
    if bs is None:
        bs = next((s for s in SHEETS_NERACA_UNCLASS if s in sheets), None)
        if bs is not None:
            flags.append("NERACA_UNCLASSIFIED")
    if bs is not None:
        row.update(scan_sheet(wb[bs], NERACA_LABELS))
    else:
        flags.append("NERACA_NONSTANDAR")

    lr_sheets = [s for s in sheets if s.strip()[0:2] in PREFIX_LABARUGI]
    if lr_sheets:
        res = {k: None for k in LABARUGI_LABELS}
        res["laba_bersih_prior"] = None
        for s in lr_sheets:
            got = scan_sheet(wb[s], LABARUGI_LABELS)
            for k, v in got.items():
                if res.get(k) is None and v is not None:
                    res[k] = v
            if res["laba_bersih_prior"] is None:
                _, prior = scan_two_numbers(wb[s], "Jumlah laba (rugi)",
                                            guards=("sebelum", "operasi", "per saham"))
                res["laba_bersih_prior"] = prior
        row.update(res)
    else:
        flags.append("TANPA_SHEET_LABARUGI")
    wb.close()

    _derive(row, flags)
    row["flag_masalah"] = ";".join(flags)
    return row


def _derive(row, flags):
    def g(k):
        v = row.get(k)
        return v if isinstance(v, float) else None
    ta, tl = g("total_aset"), g("total_liab")
    ca, cl = g("aset_lancar"), g("liab_lancar")
    eq = g("total_ekuitas")
    re_total = None
    if g("saldo_laba_ditentukan") is not None or g("saldo_laba_belum") is not None:
        re_total = (g("saldo_laba_ditentukan") or 0) + (g("saldo_laba_belum") or 0)
    lsp = g("laba_sebelum_pajak")
    bunga = g("beban_bunga") if g("beban_bunga") is not None else g("beban_keuangan_alt")
    ebit = (lsp + (bunga or 0)) if lsp is not None else None
    row["saldo_laba"] = re_total
    row["working_capital"] = (ca - cl) if (ca is not None and cl is not None) else None
    row["ebit"] = ebit
    if ta and ta != 0:
        x1 = row["working_capital"] / ta if row["working_capital"] is not None else None
        x2 = re_total / ta if re_total is not None else None
        x3 = ebit / ta if ebit is not None else None
        x4 = eq / tl if (eq is not None and tl not in (None, 0)) else None
        row["X1"], row["X2"], row["X3"], row["X4"] = x1, x2, x3, x4
        if None not in (x1, x2, x3, x4):
            row["Z2_score"] = 6.56 * x1 + 3.26 * x2 + 6.72 * x3 + 1.05 * x4
        else:
            flags.append("Z_TIDAK_LENGKAP")
        row["ROA"] = g("laba_bersih") / ta if g("laba_bersih") is not None else None
        row["leverage"] = tl / ta if tl is not None else None
        row["current_ratio"] = ca / cl if (ca is not None and cl not in (None, 0)) else None
    else:
        flags.append("TOTAL_ASET_KOSONG")


COLS = ["kode", "tahun", "nama", "industri_utama", "sektor", "subsektor",
        "mata_uang", "pembulatan", "opini_audit",
        "aset_lancar", "liab_lancar", "total_aset", "total_liab",
        "saldo_laba", "total_ekuitas", "pendapatan", "laba_sebelum_pajak",
        "beban_bunga", "laba_bersih", "laba_bersih_prior",
        "working_capital", "ebit", "X1", "X2", "X3", "X4", "Z2_score",
        "ROA", "leverage", "current_ratio", "flag_masalah"]


def main():
    rows = []
    for year in YEARS:
        folder = IN_DIR / str(year)
        if not folder.exists():
            continue
        files = sorted(folder.glob("*.xlsx"))
        print(f"=== {year}: {len(files)} files ===")
        for i, f in enumerate(files, 1):
            rows.append(parse_file(f, year))
            if i % 100 == 0:
                print(f"  ... {i}/{len(files)}")
    with open(OUT, "w", newline="", encoding="utf-8-sig") as fo:
        w = csv.DictWriter(fo, fieldnames=COLS, extrasaction="ignore")
        w.writeheader(); w.writerows(rows)
    print(f"[i] {len(rows)} rows -> {OUT}")


if __name__ == "__main__":
    main()
